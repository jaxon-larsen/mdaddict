"""PPTX paragraphs in source order, including alternate Office text and math.

This adapter selects one supported AlternateContent branch in memory. It never
writes the presentation and does not confuse compatibility fallbacks with extra
content. Pandoc still owns equation translation.
"""
from __future__ import annotations

from html import escape
from mdaddict.strategy import ExtractionIssue, SourceBlock, SourceLocator, figure, table_markdown

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
SUPPORTED = {A, M, P, "http://schemas.microsoft.com/office/drawing/2010/main"}


class MarkdownBuilder:
    def __init__(self, engine):
        self.engine = engine
        self.parts = []
        self.length = 0

    def add(self, text, block_id=None, kind="paragraph", source=None, method="native"):
        # Tesseract on Windows and model responses can contain CRLF. Normalize
        # before counting spans, rather than shifting them in final cleanup.
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        if not text:
            return
        if self.parts:
            self.parts.append("\n\n")
            self.length += 2
        start = self.length
        self.parts.append(text)
        self.length += len(text)
        if block_id:
            self.engine.source_blocks.append(SourceBlock(block_id, kind, source, start, self.length, method))

    def render(self):
        return "".join(self.parts)


def issue(engine, code, message, source, block_id=None, severity="unresolved"):
    engine.warn(message)
    engine.issues.append(ExtractionIssue(code, message, source, [block_id] if block_id else [], severity))


def select_alternates(engine, root, slide):
    """Choose a branch by namespace capability, including nested alternatives."""
    for node in list(root):
        if node.tag != f"{{{MC}}}AlternateContent":
            select_alternates(engine, node, slide)
            continue
        choices = [c for c in node if c.tag == f"{{{MC}}}Choice"]
        chosen = next((c for c in choices if c.get("Requires") and
            all(c.nsmap.get(prefix) in SUPPORTED for prefix in c.get("Requires").split())), None)
        if chosen is None:
            chosen = node.find(f"{{{MC}}}Fallback")
            issue(engine, "alternate_content_fallback",
                f"slide {slide}: unsupported alternate Office content; fallback requires review",
                SourceLocator("slide", slide))
        position = root.index(node)
        if chosen is not None:
            select_alternates(engine, chosen, slide)
            for child in list(chosen):
                root.insert(position, child)
                position += 1
        root.remove(node)


def paragraph_text(paragraph, equations):
    """Walk runs, fields, breaks and math together; never recover by text dedup."""
    def visit(node):
        if node.tag == f"{{{M}}}oMath":
            return equations.get(node, "[Office equation unresolved]")
        if node.tag == f"{{{A}}}br":
            return "<br>"
        if node.tag in {f"{{{A}}}r", f"{{{A}}}fld"}:
            text = "".join(child.text or "" for child in node if child.tag == f"{{{A}}}t")
            props = node.find(f"{{{A}}}rPr")
            baseline = int(props.get("baseline", "0")) if props is not None else 0
            return f"<sup>{text}</sup>" if baseline > 0 else f"<sub>{text}</sub>" if baseline < 0 else text
        if node.tag == f"{{{A}}}t":
            return node.text or ""
        # Property trees may contain labels that are not paragraph content.
        if node.tag in {f"{{{A}}}pPr", f"{{{A}}}endParaRPr", f"{{{M}}}oMathParaPr"}:
            return ""
        return "".join(visit(child) for child in node)

    text = visit(paragraph).strip()
    return text


def merged_table_html(table, rows):
    """GFM cannot encode spans; HTML retains actual source merge relationships."""
    output = ["<table>"]
    for row, values in zip(table.rows, rows):
        output.append("<tr>")
        for cell, value in zip(row.cells, values):
            if cell.is_spanned:
                continue
            attrs = "".join(f' {name}="{span}"' for name, span in (
                ("rowspan", cell.span_height), ("colspan", cell.span_width)) if span > 1)
            output.append(f"<td{attrs}>" + escape(value).replace("\n", "<br>") + "</td>")
        output.append("</tr>")
    return "\n".join(output + ["</table>"])


def convert_pptx(engine, path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    deck = Presentation(path)
    engine.check_pages(len(deck.slides))
    for idx, slide in enumerate(deck.slides, 1):
        select_alternates(engine, slide._element, idx)
    equations = engine._office_math(deck)
    builder = MarkdownBuilder(engine)
    consumed_paragraphs = set()

    def emit_paragraph(paragraph, source, block_id, heading=False):
        consumed_paragraphs.add(paragraph)
        text = paragraph_text(paragraph, equations)
        if not text:
            return
        props = paragraph.find(f"{{{A}}}pPr")
        level = int(props.get("lvl", "0")) if props is not None else 0
        has_math = bool(paragraph.xpath(".//*[local-name()='oMath']"))
        only_math = has_math and not "".join(paragraph.xpath(".//a:t/text()")).strip()
        kind = "heading" if heading else "equation" if only_math else "paragraph"
        if heading:
            text = "## " + text
        elif only_math and text.startswith("$") and text.endswith("$") and text.count("$") == 2:
            text = "$$\n" + text[1:-1] + "\n$$"
        elif props is None or props.find(f"{{{A}}}buNone") is None:
            text = "  " * level + "- " + text
            kind = "list_item"
        builder.add(text, block_id, kind, source)
        if "[Office equation" in text:
            issue(engine, "equation_unresolved", f"slide {source.index}, {source.element}: equation unresolved", source, block_id)

    def emit_shapes(shapes, slide_num, parent=""):
        for shape in shapes:
            element = f"{parent}shape:{shape.shape_id}"
            block_id = f"slide:{slide_num}/{element}"
            source = SourceLocator("slide", slide_num, element)
            location = f"slide {slide_num}, shape {shape.shape_id}"
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                emit_shapes(shape.shapes, slide_num, element + "/")
            elif shape.has_table:
                rows = []
                for row in shape.table.rows:
                    values = []
                    for cell in row.cells:
                        paras = [p._p for p in cell.text_frame.paragraphs]
                        consumed_paragraphs.update(paras)
                        values.append("\n".join(paragraph_text(p, equations) for p in paras))
                    rows.append(values)
                engine.check_cells(sum(map(len, rows)))
                merged = any(c.is_merge_origin or c.is_spanned for row in shape.table.rows for c in row.cells)
                text = merged_table_html(shape.table, rows) if merged else table_markdown(rows, engine.options.max_cells)
                builder.add(text, block_id, "table", source)
                if "[Office equation" in text:
                    issue(engine, "equation_unresolved", f"{location}: table equation unresolved", source, block_id)
                if merged:
                    issue(engine, "table_spans", f"{location}: merged cell spans retained in HTML; grid requires review", source, block_id, "review")
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                props = shape._element.xpath(".//p:cNvPr")
                alt = props[0].get("descr", "") if props else ""
                passage = engine.image_evidence.describe(engine, shape.image.blob, alt, location, source, block_id,
                    crop=(shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom),
                    context="\n".join(deck.slides[slide_num-1]._element.xpath(".//a:t/text()")))
                builder.add(passage.text, block_id, "figure", source, passage.method)
                for message in passage.messages:
                    issue(engine, passage.code, message, source, block_id,
                          "unresolved" if passage.method == "native" else "review")
            elif shape.has_text_frame:
                is_title = shape.is_placeholder and shape.placeholder_format.type in (1, 3)
                emitted_before = len(engine.source_blocks)
                for number, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                    paragraph_source = SourceLocator("slide", slide_num, f"{element}/paragraph:{number}")
                    emit_paragraph(paragraph._p, paragraph_source, f"{block_id}/paragraph:{number}", is_title)
                # Text on a shape does not establish the meaning of its geometry.
                vector = shape.shape_type not in (MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER)
                if vector or (len(engine.source_blocks) == emitted_before and not shape.is_placeholder):
                    figure_id = block_id + "/visual"
                    asset = engine.image_evidence.register(shape._element.xml.encode(), source, figure_id, f"{location}: shape geometry")
                    builder.add(f"![{asset.description}](asset:{asset.asset_id})", figure_id, "figure", source)
                    issue(engine, "shape_geometry_unresolved", f"{location}: shape geometry requires rendered-slide review", source, figure_id)
            else:
                asset = engine.image_evidence.register(shape._element.xml.encode(), source, block_id, f"{location}: visual object")
                builder.add(f"![{asset.description}](asset:{asset.asset_id})", block_id, "figure", source)
                issue(engine, "object_unresolved", f"{location}: chart/vector/embedded object needs rendered-slide review", source, block_id)

    for idx, slide in enumerate(deck.slides, 1):
        builder.add(f"<!-- slide: {idx} -->")
        from mdaddict.pptx_order import order_shapes
        ordered, order_problem = order_shapes(slide.shapes)
        first_block = len(engine.source_blocks)
        emit_shapes(ordered, idx)
        if order_problem:
            bids = [b.id for b in engine.source_blocks[first_block:]]
            engine.issues.append(ExtractionIssue("pptx_reading_order_review", order_problem,
                SourceLocator("slide", idx), bids, "review"))
        # Preserve unusual paragraph containers without pretending their placement is known.
        for number, paragraph in enumerate(slide._element.xpath(".//a:p"), 1):
            if paragraph not in consumed_paragraphs:
                source = SourceLocator("slide", idx, f"unplaced-paragraph:{number}")
                block_id = f"slide:{idx}/{source.element}"
                before = len(engine.source_blocks)
                emit_paragraph(paragraph, source, block_id)
                if len(engine.source_blocks) > before:
                    issue(engine, "paragraph_placement_review", f"slide {idx}: paragraph outside supported shape structure", source, block_id, "review")
        for number, node in enumerate(slide._element.xpath(".//*[local-name()='oMath']"), 1):
            if not any(parent in consumed_paragraphs for parent in node.iterancestors()):
                source = SourceLocator("slide", idx, f"unplaced-equation:{number}")
                block_id = f"slide:{idx}/{source.element}"
                builder.add(equations.get(node, "[Office equation unresolved]"), block_id, "equation", source)
                issue(engine, "equation_placement_unresolved", f"slide {idx}: equation outside supported paragraph structure", source, block_id)
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None and tf.text.strip():
                builder.add("### Speaker notes\n\n" + tf.text, f"slide:{idx}/notes", "note", SourceLocator("slide", idx, "notes"))
    return builder.render()
