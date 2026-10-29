"""Native parsers with optional Tesseract OCR. Never contacts a network service."""
from __future__ import annotations

import csv
from collections import Counter
from html.parser import HTMLParser
import io
import re
import shutil
import subprocess
import tempfile
import unicodedata
from pathlib import Path

from mdaddict.extract import IMAGE_TYPES, TEXT_TYPES
from mdaddict.strategy import (ExtractOptions, ExtractResult, ExtractionIssue, SourceBlock,
    SourceLocator, comment, figure, normalize, table_markdown)


def image_png(data: bytes) -> bytes:
    from PIL import Image, ImageOps
    with Image.open(io.BytesIO(data)) as source:
        if source.width * source.height > 40_000_000:
            raise ValueError("Image exceeds 40 megapixels")
        frame = ImageOps.exif_transpose(source).convert("RGB")
        frame.thumbnail((2400, 2400))
        output = io.BytesIO()
        frame.save(output, format="PNG")
        return output.getvalue()


def render_page(page, dpi: int) -> bytes:
    import pymupdf
    scale = min(dpi / 72, 2400 / max(page.rect.width, page.rect.height))
    return page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False).tobytes("png")


def missing_source_blocks(page, markdown: str) -> list[str]:
    """Detect dropped native glyphs without pretending to reconstruct 2-D math."""
    def significant(char):
        return char.isalnum() or unicodedata.category(char) == "Sm"
    available = Counter(char for char in markdown if significant(char))
    missing = []
    for block in page.get_text("blocks", sort=True):
        if block[6] != 0:
            continue
        text = block[4].strip()
        chars = Counter(char for char in text if significant(char))
        omitted = chars - available
        absent = sum(omitted.values())
        critical = any(unicodedata.category(c) == "Sm" or 0x1D400 <= ord(c) <= 0x1D7FF for c in omitted)
        if absent and (critical or absent >= 3 or absent / max(1, sum(chars.values())) > 0.15):
            missing.append(text)
        available.subtract(chars)
        available = +available
    return missing


def missing_decimal_tokens(native: str, markdown: str) -> list[str]:
    """Flag decimal groups lost/repeated incorrectly by native conversion.

    Character inventories cannot distinguish 1.1 from 11 plus a moved dot.
    Compare token multiplicity before recovery excerpts can mask the defect.
    This is a review guard, not numeric equivalence or a formula verifier;
    alternate representations may require review and split native tokens can
    escape it. Multi-part version numbers are deliberately excluded.
    """
    pattern = r"(?<![\d.])\d+\.\d+(?!\d|\.\d)"
    missing = Counter(re.findall(pattern, native)) - Counter(re.findall(pattern, markdown))
    return sorted(missing.elements())


def pdf_math_present(text: str) -> bool:
    return bool(re.search(r"[=∑∫√∀∃⇒⇔⊕¬∧∨]|[⁰¹²³⁴⁵⁶⁷⁸⁹]", text)) or any(
        unicodedata.category(c) == "Sm" or 0x1D400 <= ord(c) <= 0x1D7FF for c in text)


def native_code(page, markdown: str, recovered=None) -> str:
    """Recover contiguous native monospaced code using source tokens and geometry.

    Ordered matches can replace the native Markdown span. If reliable source
    code was dropped or annotations intervene, retain the surrounding text and
    append a labeled complete source excerpt rather than erase annotations.
    """
    lines = []
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = line.get("spans", [])
            if spans and all(re.search(r"mono|courier|nimbusmon|consolas", span.get("font", ""), re.I)
                             for span in spans if span.get("text", "").strip()):
                lines.append((line["bbox"], "".join(span.get("text", "") for span in spans), spans[0].get("size", 10)))
    groups = []
    for line in lines:
        if not groups or line[0][1] - groups[-1][-1][0][3] > 2 * line[2] or line[0][1] < groups[-1][-1][0][1]:
            groups.append([])
        groups[-1].append(line)
    md_lines = markdown.splitlines(keepends=True)
    recovered_groups = []
    key = lambda value: re.sub(r"[\s*_`]+", "", value)
    for group in reversed(groups):
        if len(group) < 2 or not any(re.match(r"\s*(?:for |def |class |import |from |if |while |return |public |private |protected |function |\w+\s*=)", text) for _, text, _ in group):
            continue
        matches, cursor = [], 0
        for _, text, _ in group:
            match = next((i for i in range(cursor, len(md_lines)) if key(md_lines[i]) == key(text)), None)
            if match is not None:
                matches.append(match)
                cursor = match + 1
        left = min(bounds[0] for bounds, _, _ in group)
        code = []
        for bounds, text, size in group:
            indent = max(0, round((bounds[0] - left) / (size * 0.6)))
            code.append(" " * indent + text)
        fence = "`" * max(3, max((len(m.group()) + 1 for m in re.finditer(r"`+", "\n".join(code))), default=3))
        replacement = fence + "\n" + "\n".join(code) + "\n" + fence + "\n"
        if len(matches) != len(group) or any(md_lines[i].strip() for i in range(matches[0], matches[-1] + 1) if i not in matches):
            recovered_groups.append(replacement)
            if recovered is not None:
                recovered.append(group)
        else:
            start, end = matches[0], matches[-1] + 1
            before, after = start - 1, end
            while before >= 0 and not md_lines[before].strip():
                before -= 1
            while after < len(md_lines) and not md_lines[after].strip():
                after += 1
            opening = re.match(r"^(`{3,}|~{3,})[^`~]*$", md_lines[before].strip()) if before >= 0 else None
            if opening and after < len(md_lines) and md_lines[after].strip() == opening.group(1):
                start, end = before, after + 1
            md_lines[start:end] = [replacement]
    output = "".join(md_lines)
    if recovered_groups:
        output += "\n\n### Code recovered from native layer\n\n" + "\n".join(reversed(recovered_groups))
    return output


def native_tables(markdown: str, check_cells, warn, get_found_tables) -> str:
    """Use actual table cell geometry instead of fill-forward of merged cells.

    get_found_tables is a lazy, memoized find_tables() call shared with
    analyze_page's region detection -- called at most once per page, and only
    when markdown actually looks like it contains a table.
    """
    blocks = list(re.finditer(r"(?m)^\|[^\n]*(?:\n\|[^\n]*)+", markdown))
    if not blocks:
        return markdown
    found_tables = get_found_tables()
    if found_tables is None:
        warn("PDF table mapping ambiguous; native Markdown table retained for review")
        return markdown
    tables = sorted(found_tables.tables, key=lambda table: (table.bbox[1], table.bbox[0]))
    if len(tables) != len(blocks):
        warn("PDF table mapping ambiguous; native Markdown table retained for review")
        return markdown
    replacements = []
    for table in tables:
        rows = table.extract()
        check_cells(sum(map(len, rows)))
        if table.header.external:
            rows.insert(0, table.header.names)
        if any(value is None for row in rows for value in row):
            warn("PDF merged/empty cells preserved as blank grid positions; Markdown cannot encode spans")
        replacements.append(table_markdown([["" if value is None else value for value in row] for row in rows]))
    for block, replacement in reversed(list(zip(blocks, replacements))):
        markdown = markdown[:block.start()] + replacement + markdown[block.end():]
    return markdown


class OfflineExtractor:
    name = "offline"

    def __init__(self, options: ExtractOptions | None = None):
        self.options = options or ExtractOptions()
        self.warnings: list[str] = []
        self.pages = 1
        self._cells = 0

    def warn(self, message: str):
        if message not in self.warnings:
            self.warnings.append(message)

    def check_pages(self, count: int):
        if self.options.max_pages is not None and count > self.options.max_pages:
            raise ValueError(f"Document has {count} pages/slides/sheets; limit is {self.options.max_pages}")
        self.pages = count

    def check_cells(self, count: int):
        self._cells += max(1, count)
        if self._cells > self.options.max_cells:
            raise ValueError(f"Table content exceeds --max-cells {self.options.max_cells}")

    def ocr_image(self, data: bytes, location: str) -> str:
        if not shutil.which("tesseract"):
            raise RuntimeError(f"{location} needs OCR. Install Tesseract and language data, or set ocr='off' to retain an unresolved marker.")
        with tempfile.TemporaryDirectory(prefix="mdaddict-ocr-") as directory:
            path = Path(directory) / "image.png"
            path.write_bytes(data)
            run = subprocess.run(["tesseract", str(path), "stdout", "-l", self.options.language],
                capture_output=True, timeout=self.options.timeout)
        if run.returncode:
            raise RuntimeError(f"OCR failed at {location}: {run.stderr.decode(errors='replace')[-500:]}")
        text = run.stdout.decode("utf-8").strip()
        self.warn(f"{location}: OCR transcription; equations, table structure and diagrams require review")
        return text or figure(f"{location}: no text recognized")

    def describe_image(self, data: bytes | None, alt: str, location: str) -> str:
        # Embedded pictures may be illustrations rather than text. OCR only when explicitly forced.
        if data and self.options.ocr == "force":
            return self.ocr_image(image_png(data), location)
        self.warn(f"{location}: image retained as description/marker; visual content not transcribed")
        return figure(alt or f"{location}: unresolved image")

    def _text(self, path: Path) -> str:
        try:
            return path.read_text(encoding=self.options.encoding)
        except UnicodeDecodeError as error:
            raise ValueError("Cannot decode text; set --encoding explicitly (no replacement characters inserted)") from error

    def _csv(self, path: Path) -> str:
        rows = []
        with path.open(encoding=self.options.encoding, newline="") as stream:
            for row in csv.reader(stream, delimiter="\t" if path.suffix.lower() in {".tsv", ".tab"} else ","):
                self.check_cells(len(row))
                rows.append(row)
        return table_markdown(rows, self.options.max_cells)

    def _xlsx(self, path: Path) -> str:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            self.check_pages(len(book.worksheets))
            blocks = []
            for idx, sheet in enumerate(book.worksheets, 1):
                # Avoid trusting possibly incorrect worksheet dimensions; stream actual rows.
                sheet.reset_dimensions()
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    self.check_cells(len(row))
                    rows.append(["" if value is None else str(value) for value in row])
                blocks.append(f"<!-- sheet: {idx} -->\n\n## {sheet.title}\n\n{table_markdown(rows, self.options.max_cells)}")
            self.warn("XLSX: formulas preserved as expressions, not evaluated; charts/styles/merged-cell layout are not reproduced")
            return "\n\n".join(blocks)
        finally:
            book.close()

    def _pandoc(self, path: Path) -> str:
        if path.suffix.lower() == ".docx":
            from mdaddict.docx_native import convert_docx
            return convert_docx(self, path)
        if not shutil.which("pandoc"):
            raise RuntimeError("HTML needs mdaddict[html] for ordinary markup, or Pandoc for complex markup and the compatibility fallback.")
        args = ["pandoc", "--sandbox", "--wrap=none", "-f", "html", "-t", "gfm+tex_math_dollars"]
        run = subprocess.run(args, input=self._text(path), encoding="utf-8", capture_output=True, timeout=self.options.timeout)
        if run.returncode:
            raise RuntimeError(f"Pandoc failed: {run.stderr[-500:]}")
        if run.stderr.strip():
            self.warn(f"Pandoc: {run.stderr.strip()[:500]}")
        def replace_image(match):
            alt, target = match.group(1), match.group(2).strip("<>")
            # External linked resources are never fetched.
            return self.describe_image(None, alt, f"image {target}")
        body = re.sub(r"!\[([^\]]*)\]\(([^)]+)\)(?:\{[^}]*\})?", replace_image, run.stdout)
        def replace_html_image(match):
            class Attributes(HTMLParser):
                attrs = {}
                def handle_starttag(self, tag, attrs):
                    self.attrs = dict(attrs)
            parser = Attributes()
            parser.feed(match.group())
            target = parser.attrs.get("src", "")
            return self.describe_image(None, parser.attrs.get("alt", ""), f"image {target}")
        body = re.sub(r"<img\b[^>]*>", replace_html_image, body, flags=re.I)
        return body

    def _office_math(self, deck) -> dict:
        """Convert supported OMML locally; preserve the existing fallback for others."""
        from mdaddict.office_math import convert_equation
        converted, pending = {}, []
        for idx, slide in enumerate(deck.slides, 1):
            for node in slide._element.xpath(".//*[local-name()='oMath']"):
                try:
                    converted[node] = convert_equation(node)
                except (ValueError, TypeError, RecursionError) as error:
                    pending.append((idx, node))
                    self.warn(f"slide {idx}: Office math needs fallback: {error}")
        if pending:
            if shutil.which("pandoc"):
                converted.update(self._office_math_pandoc(pending))
            else:
                self.warn("PPTX Office math unresolved: unsupported structure; install Pandoc for fallback")
                converted.update({node: "[Office equation unresolved]" for _, node in pending})
        return converted

    def _office_math_pandoc(self, entries) -> dict:
        from mdaddict.office_math import convert_with_pandoc
        return convert_with_pandoc(self, entries)

    def _pptx(self, path: Path) -> str:
        from mdaddict.pptx_native import convert_pptx
        return convert_pptx(self, path)

    def _resolve_pdf_backend(self) -> str:
        from importlib.util import find_spec
        choice = self.options.pdf_backend
        has_pdfium = find_spec("pypdfium2") is not None
        has_pymupdf = find_spec("pymupdf") is not None and find_spec("pymupdf4llm") is not None
        if choice == "pdfium":
            if not has_pdfium:
                raise RuntimeError("pdf_backend=pdfium requires pypdfium2; install mdaddict[pdf]")
            return "pdfium"
        if choice == "pymupdf":
            if not has_pymupdf:
                raise RuntimeError("pdf_backend=pymupdf requires pymupdf and pymupdf4llm; install mdaddict[pdf-pymupdf]")
            return "pymupdf"
        # auto: permissive pdfium when that is the only stack; if both extras are
        # installed (common during transition), keep PyMuPDF until callers opt in.
        if has_pdfium and not has_pymupdf:
            return "pdfium"
        if has_pymupdf:
            return "pymupdf"
        if has_pdfium:
            return "pdfium"
        raise RuntimeError("PDF support requires pypdfium2 (mdaddict[pdf]) or pymupdf (mdaddict[pdf-pymupdf])")

    def _pdf(self, path: Path) -> str:
        if self._resolve_pdf_backend() == "pdfium":
            from mdaddict.pdf_pdfium import convert_pdf
            return convert_pdf(self, path)
        import pymupdf
        import pymupdf4llm
        from mdaddict.pdf_forms import load_forms, form_markdown
        forms = load_forms(self, path)
        # Pinned model-free adapter; no layout model or implicit OCR.
        blocks = []
        offset = 0
        with pymupdf.open(path) as doc:
            if doc.needs_pass:
                raise ValueError("Encrypted PDF: provide an unlocked copy")
            self.check_pages(len(doc))
            headers = pymupdf4llm.IdentifyHeaders(doc)
            for idx, page in enumerate(doc, 1):
                recovered_math = []
                source = SourceLocator("page", idx)
                block_id = f"page:{idx}"
                method = "native"
                def issue(code, message, severity="review", element=""):
                    message = f"page {idx}: {message}"
                    self.warn(message)
                    self.issues.append(ExtractionIssue(code, message,
                        SourceLocator("page", idx, element), [block_id], severity))
                raw = page.get_text()
                real_images = [info for info in page.get_image_info()
                               if not (info.get("width", 0) <= 1 or info.get("height", 0) <= 1
                                       or min(info["bbox"][2] - info["bbox"][0], info["bbox"][3] - info["bbox"][1]) <= 3)]
                images = bool(real_images)
                # find_tables() is the single most expensive per-page call (it builds
                # a full character map to assign cell contents) and native_tables()
                # and analyze_page() both need the same result -- lazy + memoized so
                # it runs at most once, and never on a page that turns out not to
                # need it (most pages have no markdown table and no vector drawings).
                found_tables_cache = {}
                def get_found_tables():
                    if "v" not in found_tables_cache:
                        try:
                            found_tables_cache["v"] = page.find_tables(strategy="lines_strict")
                        except (AttributeError, ValueError, RuntimeError):
                            found_tables_cache["v"] = None
                    return found_tables_cache["v"]
                if self.options.ocr == "force" or (len(raw.strip()) < 20 and images):
                    if self.options.ocr == "off":
                        text = figure(f"page {idx}: scan requires OCR")
                        issue("pdf_scan_unresolved", "scan unresolved (--ocr off)", "unresolved")
                    else:
                        method = "ocr"
                        text = self.ocr_image(render_page(page, max(self.options.dpi, 200)), f"page {idx}")
                        issue("pdf_ocr_review", "OCR transcription requires source review")
                else:
                    from mdaddict.native_math import prepare_math_page, AmbiguousMath
                    math_page = prepare_math_page(page)
                    # Check omissions before appending recovered code. Repeated
                    # code characters must not mask a missing annotation in the
                    # legacy page-wide character counter.
                    try:
                        for prepared in ((math_page, None) if math_page else (None,)):
                            conversion_page = prepared.page if prepared else page
                            text = pymupdf4llm.to_markdown(prepared.doc if prepared else doc,
                                pages=[0 if prepared else idx-1], hdr_info=headers,
                                show_progress=False, write_images=False, embed_images=False, force_text=True)
                            text = native_tables(text, self.check_cells, self.warn, get_found_tables)
                            if not text.strip() and raw.strip():
                                text = conversion_page.get_text()
                                issue("pdf_plain_text_recovery", "Markdown extraction empty; preserved plain text")
                            missing = missing_source_blocks(conversion_page, text)
                            decimal_losses = missing_decimal_tokens(conversion_page.get_text(), text)
                            code_recoveries = []
                            text = native_code(conversion_page, text, code_recoveries)
                            if prepared:
                                try:
                                    text = prepared.restore(text)
                                except AmbiguousMath:
                                    continue  # Discard the copy and preserve original conversion.
                                recovered_math = prepared.expressions
                            break
                    finally:
                        if math_page:
                            math_page.doc.close()
                    if code_recoveries:
                        issue("pdf_native_code_recovery", "complete native code preserved separately; neighboring annotations retained for review")
                    if decimal_losses:
                        issue("pdf_numeric_unresolved", "native decimal tokens not preserved by Markdown conversion: "
                              + ", ".join(decimal_losses[:5]) + "; verify numeric meaning against the source", "unresolved")
                    if missing:
                        text += "\n\n### Source text recovered from native layer\n\n" + "\n\n".join(missing)
                        issue("pdf_native_recovery", f"recovered {len(missing)} native text blocks omitted by Markdown conversion; duplicates and 2-D math need review")
                    if pdf_math_present(raw):
                        issue("pdf_math_review", "mathematical notation present; native PDF extraction cannot guarantee 2-D formula layout")
                        # Raised spans are evidence of layout-dependent content, not a license to invent LaTeX.
                        for native_block in page.get_text("dict")["blocks"]:
                            for line in native_block.get("lines", []):
                                if any(span["flags"] & 1 for span in line["spans"]):
                                    bbox = ",".join(f"{v:.2f}" for v in line["bbox"])
                                    issue("pdf_math_layout_review", "raised text requires verification of exponent/annotation relationships",
                                        "review", f"bbox:{bbox}")
                    if images:
                        text += "\n\n" + figure(f"page {idx}: embedded visual content not transcribed")
                        issue("pdf_figure_unresolved", "embedded images may contain additional information", "unresolved")
                    if not text.strip():
                        text = "<!-- blank page or unresolved vector-only content -->"
                        if page.get_drawings():
                            issue("pdf_vector_unresolved", "vector-only content needs vision review", "unresolved")
                from mdaddict.fidelity_geometry import analyze_page
                analysis = analyze_page(page, text, get_found_tables, recovered_math=recovered_math)
                text = analysis.markdown
                if method == "native":
                    from mdaddict.pdf_order import reorder_numbered_columns
                    ordered_text, order_reason = reorder_numbered_columns(page, text)
                    if order_reason:
                        issue("pdf_reading_order_unresolved", order_reason, "unresolved")
                    elif ordered_text != text:
                        text = ordered_text
                        issue("pdf_reading_order_recovery", "numbering and disjoint text geometry corroborate paragraph order",
                              "informational")
                if "\ufffd" in text:
                    issue("pdf_glyph_unresolved", "unreadable glyphs; use OCR or AI", "unresolved")
                for region in analysis.regions:
                    issue("pdf_region_unresolved", region.reason, "unresolved",
                          "bbox:" + ",".join(f"{v:.2f}" for v in region.bbox))
                text += form_markdown(forms, idx)
                chunk = normalize(f"<!-- page: {idx} -->\n\n{text.strip()}")
                self.source_blocks.append(SourceBlock(block_id, "page", source, offset, offset + len(chunk), method))
                offset += len(chunk) + 2
                blocks.append(chunk)
        return "\n\n".join(blocks)

    def _image(self, path: Path) -> str:
        if path.suffix.lower() in {".heic", ".heif"}:
            from pillow_heif import register_heif_opener
            register_heif_opener()
        from PIL import Image
        blocks = []
        offset = 0
        with Image.open(path) as source:
            count = getattr(source, "n_frames", 1)
            self.check_pages(count)
            for idx in range(count):
                source.seek(idx)
                if source.width * source.height > 40_000_000:
                    raise ValueError("Image exceeds 40 megapixels")
                stream = io.BytesIO()
                source.save(stream, format="PNG")
                data = image_png(stream.getvalue())
                if self.options.ocr == "off":
                    text = figure(f"frame {idx+1}: OCR disabled")
                    self.warn("Image content unresolved (--ocr off)")
                else:
                    text = self.ocr_image(data, f"frame {idx+1}")
                chunk = f"<!-- frame: {idx+1} -->\n\n{text}"
                block_id = f"frame:{idx+1}"
                locator = SourceLocator("frame", idx+1)
                self.source_blocks.append(SourceBlock(block_id, "image", locator, offset, offset + len(chunk), "ocr" if self.options.ocr != "off" else "native"))
                unresolved = self.options.ocr == "off" or not text.strip() or text.lstrip().startswith("<!-- fig:")
                self.issues.append(ExtractionIssue("image_text_unavailable" if unresolved else "image_transcription_review",
                    "Image text was not transcribed." if unresolved else "OCR text requires review; diagrams and spatial relationships are not certified.",
                    locator, [block_id], "unresolved" if unresolved else "review"))
                blocks.append(chunk)
                offset += len(chunk) + 2
        return "\n\n".join(blocks)

    def extract(self, path: Path) -> ExtractResult:
        from mdaddict.image_evidence import ImageEvidence
        self.image_evidence = ImageEvidence()
        self.warnings = []
        self.source_blocks = []
        self.issues = []
        self.pdf_forms = None
        self.pages = 1
        self._cells = 0
        suffix = path.suffix.lower()
        if suffix in TEXT_TYPES:
            body, backend = self._text(path), "plain"
        elif suffix in {".csv", ".tsv", ".tab"}:
            body, backend = self._csv(path), "csv"
        elif suffix == ".xlsx":
            body, backend = self._xlsx(path), "openpyxl"
        elif suffix == ".docx":
            body, backend = self._pandoc(path), "pandoc"
        elif suffix in {".html", ".htm", ".xhtml"}:
            from mdaddict.html_native import convert_html
            body, backend = convert_html(self, path)
        elif suffix == ".pptx":
            body, backend = self._pptx(path), "python-pptx"
        elif suffix == ".pdf":
            pdf_backend = self._resolve_pdf_backend()
            body, backend = self._pdf(path), ("pypdfium2" if pdf_backend == "pdfium" else "pymupdf4llm")
        elif suffix in IMAGE_TYPES:
            body, backend = self._image(path), "tesseract"
        else:
            from mdaddict.delegated import extract_extended
            return extract_extended(path, self.options)
        if not body.strip():
            self.warn("Input produced empty Markdown")
        body = normalize(body)
        if not self.source_blocks and body:
            self.source_blocks = [SourceBlock("document", "document", SourceLocator(), 0, len(body), "mixed")]
        # Other adapters retain coarse document provenance until migrated: these
        # warnings carry no block_ids, so ExtractionIssue.__post_init__ treats
        # them as informational (a document-wide issue can never be resolved by
        # correcting a specific block).
        recorded = {issue.message for issue in self.issues}
        for warning in self.warnings:
            if warning not in recorded:
                self.issues.append(ExtractionIssue("extraction_warning", warning))
        result = ExtractResult("", body, path.name, f"{self.name}:{backend}", self.pages,
            list(self.warnings), blocks=self.source_blocks, issues=self.issues)
        if self.pdf_forms is not None:
            result.metadata["pdf_forms"] = self.pdf_forms
        result.metadata["provenance_granularity"] = "paragraph_or_object" if suffix == ".pptx" else "xml_paragraph" if suffix == ".docx" else "page" if suffix == ".pdf" and self.source_blocks and self.source_blocks[0].source.unit == "page" else "document"
        if self.image_evidence.records:
            result.metadata.update(self.image_evidence.metadata())
        result.refresh_status()
        return result
