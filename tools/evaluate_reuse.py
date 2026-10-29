"""Bounded offline reuse probe; reuses the existing source-content evaluator.

Run with the evaluation environment's Python. This is a synthetic feasibility
probe, not a general accuracy benchmark or a production converter interface.
"""
from __future__ import annotations

import argparse
from contextlib import redirect_stdout, redirect_stderr
import hashlib
from importlib import metadata
import io
import json
import os
from pathlib import Path
import socket
import statistics
import subprocess
import sys
import time


BACKENDS = {
    "mdaddict": {"html", "pdf", "docx", "pptx", "xlsx"},
    "markdownify": {"html"},
    "pandoc": {"html", "docx"},
    "mammoth-markdownify": {"docx"},
    "python-docx-text": {"docx"},
    "markitdown": {"html", "pdf", "docx", "pptx", "xlsx"},
    "docling-native": {"html", "docx", "pptx", "xlsx"},
    "pdfium-text": {"pdf"},
    "pypdf-text": {"pdf"},
    "pdfplumber-text": {"pdf"},
}


def dump(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def prepare(root, reference):
    """Expected content is fixed here before any candidate converter is run."""
    package = Path(__file__).resolve().parents[1]
    fixture = package / "tests/fixtures"
    manifest = json.loads((reference / "engine/tests/fixtures/benchmark/heldout.json").read_text())
    cases = []
    for name, fmt in (("conditions.pdf", "pdf"), ("measurement.docx", "docx"), ("comparison.docx", "docx"), ("procedures.pptx", "pptx")):
        source = fixture / name
        if not source.exists():
            source = reference / "engine/tests/fixtures/benchmark/heldout" / name
        row = next(s for s in manifest["sources"] if s["source_name"] == name)
        assert hashlib.sha256(source.read_bytes()).hexdigest() == row["source_sha256"]
        expected = "\n\n".join(check["expected"] for case in row["cases"] for check in case["checks"] if check["kind"] == "coverage")
        cases.append({"id": name, "format": fmt, "path": str(source), "expected": expected})
    synthetic = json.loads((reference / "engine/tests/fixtures/benchmark/synthetic-office.json").read_text())
    for row in synthetic["sources"]:
        if row["format"] not in {"docx", "pptx"}:
            continue
        source = reference / "engine/tests/fixtures/benchmark/synthetic" / row["source_name"]
        assert hashlib.sha256(source.read_bytes()).hexdigest() == row["source_sha256"]
        expected = "\n\n".join(check["expected"] for case in row["cases"] for check in case["checks"] if check["kind"] == "coverage")
        cases.append({"id": "algebra." + row["format"], "format": row["format"], "path": str(source), "expected": expected})
    html = [
        ("html-basic", "<h1>Measurement</h1><p>Use <strong>two</strong> trials.</p><p>Keep λ and −3.</p>", "# Measurement\n\nUse **two** trials.\n\nKeep λ and −3."),
        ("html-repetition", "<h1>Conditions</h1><p>Keep the units.</p><p>Keep the units.</p><p>A missing value is not zero.</p>", "# Conditions\n\nKeep the units.\n\nKeep the units.\n\nA missing value is not zero."),
        ("html-table", "<table><tr><th>Case</th><th>Value</th></tr><tr><td>A</td><td>3</td></tr><tr><td>B</td><td></td></tr></table>", "| Case | Value |\n|---|---|\n| A | 3 |\n| B | |"),
        ("html-spans", '<table><tr><th colspan="2">Pair</th></tr><tr><td>A</td><td>B</td></tr></table>', '<table><tr><th colspan="2">Pair</th></tr><tr><td>A</td><td>B</td></tr></table>'),
        ("html-math", '<math xmlns="http://www.w3.org/1998/Math/MathML"><mfrac><mrow><mi>x</mi><mo>+</mo><mn>1</mn></mrow><mn>2</mn></mfrac></math>', r"$\frac{x+1}{2}$"),
    ]
    for name, source, expected in html:
        path = root / "inputs" / (name + ".html")
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(source, encoding="utf-8")
        cases.append({"id": name, "format": "html", "path": str(path), "expected": expected})
    from openpyxl import Workbook
    path = root / "inputs/formulas.xlsx"
    book = Workbook()
    book.active.title = "Values"
    book.active.append(["Input", "Formula"])
    book.active.append([3, "=A2*2"])
    book.save(path)
    book.close()
    cases.append({"id": "formulas.xlsx", "format": "xlsx", "path": str(path),
                  "expected": "## Values\n\n| Input | Formula |\n|---|---|\n| 3 | =A2*2 |"})
    for case in cases:
        case["sha256"] = hashlib.sha256(Path(case["path"]).read_bytes()).hexdigest()
    dump(root / "cases.json", cases)
    return cases


def converter(name):
    if name == "mdaddict":
        from mdaddict import ExtractOptions, extract
        def run(path):
            result = extract(path, options=ExtractOptions(ocr="off"))
            return result.markdown, {"blocks": len(result.blocks), "issues": len(result.issues)}
        return run
    if name == "markdownify":
        from markdownify import markdownify
        return lambda p: (markdownify(p.read_text(encoding="utf-8"), heading_style="ATX", sub_symbol="<sub>", sup_symbol="<sup>"), {})
    if name == "pandoc":
        import pypandoc
        return lambda p: (pypandoc.convert_file(str(p), "gfm+tex_math_dollars", extra_args=["--sandbox", "--wrap=none"]), {})
    if name == "mammoth-markdownify":
        import mammoth
        from markdownify import markdownify
        def run(path):
            with path.open("rb") as stream:
                result = mammoth.convert_to_html(stream)
            return markdownify(result.value, heading_style="ATX", sub_symbol="<sub>", sup_symbol="<sup>"), {"messages": [str(m) for m in result.messages]}
        return run
    if name == "python-docx-text":
        from docx import Document
        from docx.text.paragraph import Paragraph
        def run(path):
            parts = []
            for item in Document(path).iter_inner_content():
                if isinstance(item, Paragraph):
                    parts.append(item.text)
                else:
                    parts.extend(" | ".join(c.text for c in row.cells) for row in item.rows)
            return "\n\n".join(parts), {"parser_only": True}
        return run
    if name == "markitdown":
        from markitdown import MarkItDown
        engine = MarkItDown(enable_plugins=False)
        return lambda p: (engine.convert_local(str(p)).markdown, {})
    if name == "docling-native":
        from docling.document_converter import DocumentConverter
        from docling.datamodel.base_models import InputFormat
        engine = DocumentConverter(allowed_formats=[InputFormat.HTML, InputFormat.DOCX, InputFormat.PPTX, InputFormat.XLSX])
        return lambda p: (engine.convert(str(p)).document.export_to_markdown(), {})
    if name == "pdfium-text":
        import pypdfium2
        def run(path):
            parts = []
            with pypdfium2.PdfDocument(path) as document:
                for page in document:
                    text = page.get_textpage()
                    try:
                        parts.append(text.get_text_bounded())
                    finally:
                        text.close()
                        page.close()
            return "\n\n".join(parts), {"parser_only": True}
        return run
    if name == "pypdf-text":
        from pypdf import PdfReader
        return lambda p: ("\n\n".join(page.extract_text() for page in PdfReader(p).pages), {"parser_only": True})
    if name == "pdfplumber-text":
        import pdfplumber
        def run(path):
            with pdfplumber.open(path) as document:
                return "\n\n".join(page.extract_text() or "" for page in document.pages), {"parser_only": True}
        return run
    raise ValueError(name)


def worker(args):
    def blocked(*_args, **_kwargs):
        raise RuntimeError("Network calls are forbidden in the offline reuse probe")
    socket.create_connection = blocked
    socket.socket.connect = blocked
    cases = json.loads(args.cases.read_text(encoding="utf-8"))
    started = time.perf_counter()
    with redirect_stdout(sys.stderr):
        run = converter(args.worker)
    rows = []
    import_ms = (time.perf_counter() - started) * 1000
    for case in cases:
        if case["format"] not in BACKENDS[args.worker]:
            continue
        path = Path(case["path"])
        assert hashlib.sha256(path.read_bytes()).hexdigest() == case["sha256"]
        samples = []
        try:
            first_output = None
            with redirect_stdout(sys.stderr):
                for _ in range(args.repeat + 1):
                    started = time.perf_counter()
                    markdown, details = run(path)
                    samples.append((time.perf_counter() - started) * 1000)
                    if first_output is None:
                        first_output = markdown
                    assert markdown == first_output, "Output changed between repeated conversions"
            rows.append({"case": case["id"], "first_ms": samples[0], "warm_median_ms": statistics.median(samples[1:]),
                         "markdown": markdown, "details": details})
        except Exception as error:
            rows.append({"case": case["id"], "error": f"{type(error).__name__}: {error}"})
    dump(args.result, {"backend": args.worker, "import_ms": import_ms, "rows": rows})


def main(args):
    import psutil
    import re
    args.out.mkdir(parents=True, exist_ok=True)
    cases = prepare(args.out, args.reference_root)
    sys.path.insert(0, str(args.reference_root / "engine/tools"))
    from benchmark_content import compare_content
    results = []
    for backend in BACKENDS:
        python = args.baseline_python if backend == "mdaddict" else Path(sys.executable)
        result = args.out / (backend + ".json")
        log = args.out / (backend + ".log")
        command = [str(python), str(Path(__file__).resolve()), "--worker", backend,
                   "--cases", str(args.out / "cases.json"), "--result", str(result), "--repeat", str(args.repeat)]
        env = {k: v for k, v in os.environ.items() if k not in {"PYTHONPATH", "PYTHONHOME"}}
        env["PYTHONUTF8"] = "1"
        peak = 0
        started = time.perf_counter()
        with log.open("w", encoding="utf-8") as stream:
            process = subprocess.Popen(command, cwd=args.out, env=env, stdout=stream, stderr=stream)
            observed = psutil.Process(process.pid)
            while process.poll() is None:
                try:
                    total = sum(p.memory_info().rss for p in [observed, *observed.children(recursive=True)])
                    peak = max(peak, total)
                except psutil.Error:
                    pass
                if time.perf_counter() - started > 120:
                    for child in observed.children(recursive=True):
                        child.kill()
                    process.kill()
                    raise TimeoutError(backend)
                time.sleep(.01)
        if process.returncode:
            raise RuntimeError(f"{backend} worker failed; inspect {log}")
        data = json.loads(result.read_text(encoding="utf-8"))
        data["peak_process_tree_rss_bytes"] = peak
        for row in data["rows"]:
            case = next(c for c in cases if c["id"] == row["case"])
            if "error" in row:
                row["comparison"] = {"status": "error"}
                continue
            # Strip converter-generated pagination labels, never figure markers.
            markdown = re.sub(r"<!--\s*(?:page|slide(?: number)?|sheet):[^>]*-->", "", row["markdown"], flags=re.I)
            # GFM's $`...`$ and ordinary $...$ delimit the same inline math.
            markdown = re.sub(r"\$`([^`\n]+)`\$", r"$\1$", markdown)
            try:
                row["comparison"] = compare_content(markdown, case["expected"])
            except ValueError as error:
                row["comparison"] = {"status": "unsupported_comparison", "reason": str(error)}
        results.append(data)
        print(json.dumps({"backend": backend, "import_ms": round(data["import_ms"], 2),
                          "peak_mib": round(peak / 2**20, 1),
                          "cases": [{"id": r["case"], "status": r["comparison"]["status"],
                                     "warm_ms": round(r.get("warm_median_ms", 0), 2)} for r in data["rows"]]}), flush=True)
    dump(args.out / "report.json", {"scope": "12 synthetic documents; offline; full-document content checks, not location parity",
                                    "repeat": args.repeat, "python": sys.version, "results": results})


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reference-root", type=Path)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--baseline-python", type=Path)
    parser.add_argument("--repeat", type=int, default=3)
    parser.add_argument("--worker", choices=BACKENDS)
    parser.add_argument("--cases", type=Path)
    parser.add_argument("--result", type=Path)
    args = parser.parse_args()
    if args.worker:
        worker(args)
    else:
        if not all((args.reference_root, args.out, args.baseline_python)):
            parser.error("--reference-root, --out, and --baseline-python are required")
        args.reference_root = args.reference_root.resolve()
        args.out = args.out.resolve()
        args.baseline_python = args.baseline_python.resolve()
        main(args)
