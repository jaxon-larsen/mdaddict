"""Source occurrence identity, no bitmap retention, and visual quarantine."""
from dataclasses import asdict
import hashlib
import io

import pytest
from mdaddict import AssetRecord, ExtractResult, ExtractOptions, SourceLocator, extract


def test_asset_identity_validates_occurrence_and_geometry():
    source = SourceLocator('slide', 1, 'shape:3')
    digest = hashlib.sha256(b'image').hexdigest()
    a = AssetRecord.create(source, (0, 0, 1, 1), digest, 'A points to B', ['b'])
    assert AssetRecord(**asdict(a)) == a
    assert a.asset_id == AssetRecord.create(source, (0., 0., 1., 1.), digest, 'new alt', ['b']).asset_id
    assert a.asset_id != AssetRecord.create(source, (.1, 0, 1, 1), digest, 'alt', ['b']).asset_id
    with pytest.raises(ValueError):
        AssetRecord(**{**asdict(a), 'crop': [0, 0, 0, 1]})
    with pytest.raises(ValueError):
        AssetRecord(**{**asdict(a), 'crop': [0, 0, float('nan'), 1]})
    with pytest.raises(ValueError):
        AssetRecord(**{**asdict(a), 'asset_id': 'ast_forged'})
    with pytest.raises(ValueError, match='unknown'):
        ExtractResult('hash', 'text', 'source', 'test', assets=[a])


def test_pptx_picture_assets_survive_roundtrip_without_bitmap_sidecars(tmp_path):
    Presentation = pytest.importorskip("pptx").Presentation
    from pptx.util import Inches
    Image = pytest.importorskip("PIL.Image")
    image = io.BytesIO()
    Image.new('RGB', (50, 50), 'white').save(image, format='PNG')
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(image, Inches(1), Inches(1))
    path = tmp_path / 'diagram.pptx'
    deck.save(path)
    output = tmp_path / 'out'
    result = extract(path, options=ExtractOptions(ocr='off'), out_dir=output)
    assert len(result.assets) == 1
    assert result.assets[0].visual_hash == hashlib.sha256(image.getvalue()).hexdigest()
    assert f'(asset:{result.assets[0].asset_id})' in result.markdown
    assert any(i.code == 'asset_context_unresolved' and i.block_ids == result.assets[0].block_ids for i in result.issues)
    assert result.review_status == 'needs_review'
    replay = extract(path, options=ExtractOptions(ocr='off'), out_dir=output)
    assert asdict(replay.assets[0]) == asdict(result.assets[0])
    assert replay.cache_hit
    assert not list(output.rglob('*.png'))
    for block in result.blocks:
        assert 0 <= block.start <= block.end <= len(result.markdown)


def test_pdf_vector_page_keeps_a_stable_text_first_reference(tmp_path):
    pypdf = pytest.importorskip('pypdf')
    pytest.importorskip('pypdfium2')
    pytest.importorskip('PIL.Image')
    from pypdf.generic import DecodedStreamObject, NameObject
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(200, 200)
    commands = DecodedStreamObject()
    commands.set_data(b'20 20 m 150 150 l S')
    page[NameObject('/Contents')] = writer._add_object(commands)
    source = tmp_path / 'vector.pdf'
    writer.write(source)
    writer.close()
    result = extract(source, options=ExtractOptions(pdf_backend='pdfium', ocr='off'), out_dir=tmp_path/'out')
    assert result.assets and result.assets[0].locator.index == 1
    assert result.assets[0].block_ids == ['page:1']
    assert f'(asset:{result.assets[0].asset_id})' in result.markdown
    assert result.status == 'partial'
    assert not list((tmp_path/'out').rglob('*.png'))
